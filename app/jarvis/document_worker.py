"""One-shot structured document worker (XLSX/DOCX in .venv-skills, PDF in .venv).

Reads one bounded JSON request on stdin and prints one JSON reply. It never
replaces the user's file: writes go to the host-chosen temporary output, which
the host verifies independently before an atomic commit. Not a security sandbox;
it only parses documents with maintained libraries (openpyxl, python-docx, pypdf).
"""
import json
import re
import sys

LIMIT = 1_000_000
CELL = re.compile(r'[A-Z]{1,3}[1-9][0-9]{0,6}')


def _cell_value(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)  # dates/times: ISO-like text for JSON transport


def xlsx_read(request):
    from openpyxl import load_workbook
    path, sheet = request['path'], request.get('sheet')
    formulas = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    cached = load_workbook(path, data_only=True, read_only=False, keep_links=True)
    names = formulas.sheetnames
    selected = [sheet] if sheet else names
    result = {'sheets': names, 'cells': {}, 'formula_cells': 0}
    limit = int(request.get('max_cells', 2000))
    count = 0
    for name in selected:
        if name not in names:
            raise ValueError('Unknown worksheet.')
        rows = {}
        for row in formulas[name].iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                count += 1
                if count > limit:
                    raise ValueError('Workbook exceeds the cell read limit.')
                raw = cell.value
                formula = isinstance(raw, str) and raw.startswith('=')
                result['formula_cells'] += formula
                rows[cell.coordinate] = {'value': _cell_value(cached[name][cell.coordinate].value),
                                         **({'formula': raw} if formula else {})}
        result['cells'][name] = rows
    return result


def xlsx_write(request):
    from openpyxl import load_workbook
    path, output, sheet, cells = request['path'], request['output'], request['sheet'], request['cells']
    book = load_workbook(path, data_only=False, keep_links=True)
    if sheet not in book.sheetnames:
        raise ValueError('Unknown worksheet.')
    target = book[sheet]
    for coordinate, value in cells.items():
        if not CELL.fullmatch(coordinate):
            raise ValueError('Cells must be A1-style coordinates.')
        target[coordinate] = value
    book.save(output)
    return {'written': sorted(cells), 'output': output}


def docx_read(request):
    from docx import Document
    document = Document(request['path'])
    paragraphs = [{'text': p.text, 'style': p.style.name if p.style is not None else ''} for p in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    if len(paragraphs) > 5000 or sum(len(t) for t in tables) > 5000:
        raise ValueError('Document exceeds the structure read limit.')
    return {'paragraphs': paragraphs, 'tables': tables, 'sections': len(document.sections),
            'core_title': document.core_properties.title or ''}


def docx_write(request):
    from docx import Document
    document = Document(request['path'])
    for edit in request['edits']:
        kind = edit['kind']
        if kind == 'append_paragraph':
            style = edit.get('style') or None
            if style is not None and style not in {s.name for s in document.styles}:
                raise ValueError('Unknown paragraph style.')
            document.add_paragraph(edit['text'], style=style)
        elif kind == 'replace_text':
            find, replace = edit['find'], edit['replace']
            hits = 0
            for paragraph in document.paragraphs:
                for run in paragraph.runs:
                    if find in run.text:
                        run.text = run.text.replace(find, replace)
                        hits += 1
            # Text split across runs would need reformatting; refuse rather than flatten styles.
            if hits == 0:
                raise ValueError('Text not found within a single formatted run; no change made.')
        elif kind == 'set_table_cell':
            table = document.tables[edit['table']]
            cell = table.cell(edit['row'], edit['column'])
            paragraphs = cell.paragraphs
            if len(paragraphs) != 1 or len(paragraphs[0].runs) > 1:
                raise ValueError('Table cell has mixed formatting; no change made.')
            if paragraphs[0].runs:
                paragraphs[0].runs[0].text = edit['text']
            else:
                paragraphs[0].add_run(edit['text'])
        else:
            raise ValueError('Unsupported document edit.')
    document.save(request['output'])
    return {'edits': len(request['edits']), 'output': request['output']}


def pdf_inspect(request):
    from pypdf import PdfReader
    reader = PdfReader(request['path'])
    if reader.is_encrypted:
        return {'encrypted': True, 'pages': None}
    metadata = reader.metadata or {}
    return {'encrypted': False, 'pages': len(reader.pages),
            'title': str(metadata.get('/Title', '') or ''), 'producer': str(metadata.get('/Producer', '') or '')}


def pdf_text(request):
    from pypdf import PdfReader
    reader = PdfReader(request['path'])
    if reader.is_encrypted:
        raise ValueError('Encrypted PDF; no text extracted.')
    pages = request.get('pages') or list(range(1, min(len(reader.pages), 50) + 1))
    text = {}
    for number in pages:
        if not 1 <= number <= len(reader.pages):
            raise ValueError('Page number outside the document.')
        text[str(number)] = (reader.pages[number - 1].extract_text() or '')[:20000]
    return {'pages': len(reader.pages), 'text': text}


def pdf_compose(request):
    """New output only: merge sources, extract selected pages, or rotate pages."""
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    for source in request['sources']:
        reader = PdfReader(source['path'])
        if reader.is_encrypted:
            raise ValueError('Encrypted PDF; no output written.')
        numbers = source.get('pages') or list(range(1, len(reader.pages) + 1))
        for number in numbers:
            if not 1 <= number <= len(reader.pages):
                raise ValueError('Page number outside the document.')
            page = reader.pages[number - 1]
            rotation = source.get('rotate', 0)
            if rotation:
                page.rotate(rotation)
            writer.add_page(page)
    with open(request['output'], 'xb') as handle:  # never overwrite
        writer.write(handle)
    return {'pages': len(writer.pages), 'output': request['output']}


OPERATIONS = {'xlsx_read': xlsx_read, 'xlsx_write': xlsx_write, 'docx_read': docx_read, 'docx_write': docx_write,
              'pdf_inspect': pdf_inspect, 'pdf_text': pdf_text, 'pdf_compose': pdf_compose}


def main():
    try:
        raw = sys.stdin.buffer.read(LIMIT + 1)
        if len(raw) > LIMIT:
            raise ValueError('Document request exceeds 1 MB.')
        request = json.loads(raw)
        operation = OPERATIONS[request['operation']]
        text = json.dumps({'ok': True, 'value': operation(request)}, allow_nan=False, ensure_ascii=False)
        if len(text.encode('utf-8')) > 4 * LIMIT:
            raise ValueError('Document result exceeds 4 MB.')
    except Exception as error:  # Reported to the host as a typed failure, never as success.
        text = json.dumps({'ok': False, 'error_type': type(error).__name__, 'error': str(error)[:500]})
    sys.stdout.buffer.write(text.encode('utf-8'))


if __name__ == '__main__':
    main()
