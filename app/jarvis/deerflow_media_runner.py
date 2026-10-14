"""Trusted, keyless media recipes; invoked only in the offline Docker job."""
import json
from pathlib import Path
import subprocess
import textwrap
import wave


def render(spec):
    import numpy as np
    from PIL import Image, ImageDraw, ImageFont
    out = Path('/mnt/user-data/outputs')
    kind, title = spec['kind'], spec.get('title', 'Jarvis')
    text = spec.get('text', '')
    if kind in {'chart', 'poster', 'video'}:
        if kind == 'chart':
            import matplotlib.pyplot as plt
            plt.style.use('dark_background')
            fig, ax = plt.subplots(figsize=(12.8, 7.2), dpi=100)
            ax.bar(spec['labels'], spec['values'], color='#40e0b5')
            ax.set_title(title)
            fig.tight_layout(); fig.savefig(out / 'chart.png'); plt.close(fig)
        else:
            image = Image.new('RGB', (1280, 720), '#101b2d')
            draw = ImageDraw.Draw(image)
            big = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 48)
            small = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 26)
            draw.rectangle((60, 70, 90, 650), fill='#40e0b5')
            draw.multiline_text((125, 140), textwrap.fill(title, 32), font=big, fill='white', spacing=12)
            draw.multiline_text((130, 380), textwrap.fill(text, 68), font=small, fill='#b8cfdf', spacing=10)
            image.save(out / 'poster.png')
        if kind == 'video':
            subprocess.run(['ffmpeg', '-nostdin', '-n', '-v', 'error', '-loop', '1', '-i', str(out / 'poster.png'),
                '-t', str(spec.get('duration', 6)), '-r', '24', '-c:v', 'libx264', '-preset', 'fast', '-crf', '24',
                '-threads', '2', '-pix_fmt', 'yuv420p', str(out / 'video.mp4')], check=True, timeout=40)
    elif kind == 'pdf':
        from reportlab.pdfgen import canvas
        page = canvas.Canvas(str(out / 'report.pdf'))
        page.setFont('Helvetica-Bold', 20); page.drawString(45, 790, title[:60])
        page.setFont('Helvetica', 12)
        y = 750
        for line in textwrap.wrap(text, 80):
            if y < 50:
                page.showPage(); page.setFont('Helvetica', 12); y = 790
            page.drawString(45, y, line); y -= 18
        page.save()
    elif kind == 'slides':
        from pptx import Presentation
        from pptx.util import Pt
        deck = Presentation()
        for item in spec['slides']:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = item['title']
            slide.placeholders[1].text = item['text']
            for paragraph in slide.shapes.title.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(26)
            for paragraph in slide.placeholders[1].text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(18)
        deck.save(out / 'slides.pptx')
    elif kind == 'music':
        rate, duration = 24000, int(spec.get('duration', 8))
        times = np.arange(rate * duration) / rate
        notes = np.array([261.63, 329.63, 392, 523.25, 392, 329.63, 293.66, 261.63])
        frequency = notes[np.arange(rate*duration)//rate % len(notes)]
        samples = .25 * np.sin(np.pi*(times % 1))**2 * np.sin(2*np.pi*frequency*times)
        with wave.open(str(out / 'melody.wav'), 'wb') as sound:
            sound.setnchannels(1); sound.setsampwidth(2); sound.setframerate(rate)
            sound.writeframes((samples*32767).astype(np.int16).tobytes())
    elif kind == 'narration':
        from kokoro_onnx import Kokoro
        import onnxruntime as ort
        import soundfile as sf
        settings = ort.SessionOptions(); settings.intra_op_num_threads = 2; settings.inter_op_num_threads = 1
        native = ort.InferenceSession('/mnt/models/kokoro/kokoro-v1.0.onnx', sess_options=settings, providers=['CPUExecutionProvider'])
        class FloatSpeed:
            def __getattr__(self, name):
                return getattr(native, name)
            def run(self, names, inputs):
                inputs = dict(inputs)
                inputs['speed'] = np.array([1.0], dtype=np.float32)
                return native.run(names, inputs)
        voice = Kokoro.from_session(FloatSpeed(), '/mnt/models/kokoro/voices-v1.0.bin')
        audio, rate = voice.create(text, voice=spec.get('voice', 'am_michael'), speed=1, lang='en-us')
        if not audio.size:
            raise ValueError('Speech output was empty.')
        sf.write(out / 'narration.wav', audio, rate)
    else:
        raise ValueError('Unsupported free media recipe.')
    return [{'name': p.name, 'bytes': p.stat().st_size} for p in sorted(out.iterdir()) if p.is_file()]


if __name__ == '__main__':
    print(json.dumps(render(json.loads(Path('spec.json').read_text()))))
