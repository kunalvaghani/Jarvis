"""Trusted synthetic load fixture, executed only inside the offline Docker job."""
import json
from pathlib import Path
import resource
import subprocess
import time
import wave
import zipfile

import numpy as np
import pandas as pd
import duckdb
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from sklearn.linear_model import LogisticRegression

out = Path('/mnt/user-data/outputs')
started = time.perf_counter()
steps = {}

stage = time.perf_counter()
n = 1_000_000
values = np.arange(n, dtype=np.int64)
frame = pd.DataFrame({'group': values % 100, 'value': values, 'square': values * values})
grouped = frame.groupby('group').agg({'value': 'sum', 'square': 'sum'})
expected = n * (n-1) * (2*n-1) // 6
actual = int(grouped['square'].sum())
assert actual == expected
db = duckdb.connect()
db.register('data', frame)
sql = db.execute('select count(*), sum(square)::HUGEINT from data').fetchone()
assert sql == (n, expected)
grouped.to_csv(out / 'groups.csv')
steps['million_row_analysis_seconds'] = round(time.perf_counter() - stage, 3)

stage = time.perf_counter()
rng = np.random.default_rng(42)
x = rng.normal(size=(6000, 6))
y = (x[:, 0] + .7*x[:, 1] - .4*x[:, 2] > 0).astype(int)
model = LogisticRegression(max_iter=200).fit(x[:4000], y[:4000])
accuracy = float(model.score(x[4000:], y[4000:]))
assert accuracy > .97
steps['held_out_ml_seconds'] = round(time.perf_counter() - stage, 3)

stage = time.perf_counter()
plt.style.use('dark_background')
fig, ax = plt.subplots(figsize=(12.8, 7.2), dpi=100)
fig.patch.set_facecolor('#101b2d')
ax.set_facecolor('#101b2d')
ax.bar(grouped.index, grouped['value'], color='#40e0b5', width=.85)
ax.set(title='Jarvis offline sandbox | 1,000,000-row analysis', xlabel='Synthetic group', ylabel='Sum of values')
ax.spines[['top', 'right']].set_visible(False)
fig.tight_layout()
fig.savefig(out / 'chart.png')
plt.close(fig)
with Image.open(out / 'chart.png') as image:
    assert image.size == (1280, 720)
doc = canvas.Canvas(str(out / 'report.pdf'))
doc.setTitle('Synthetic offline sandbox validation')
doc.setFont('Helvetica-Bold', 18)
doc.drawString(40, 790, 'Jarvis offline sandbox validation')
doc.setFont('Helvetica', 11)
for line, text in enumerate([f'Rows: {n:,}', f'Exact sum of squares: {actual}', f'Held-out accuracy: {accuracy:.4f}',
                           'Generated locally; no API or account credentials.']):
    doc.drawString(40, 755 - line*22, text)
doc.drawImage(str(out / 'chart.png'), 40, 380, width=515, height=290)
doc.save()
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(12.8), Inches(7.2)
for index in range(5):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(out / 'chart.png'), 0, 0, width=deck.slide_width, height=deck.slide_height)
deck.save(out / 'slides.pptx')
with zipfile.ZipFile(out / 'slides.pptx') as archive:
    assert len([name for name in archive.namelist() if name.startswith('ppt/slides/slide') and name.endswith('.xml')]) == 5
assert (out / 'report.pdf').read_bytes().startswith(b'%PDF-')
steps['chart_pdf_slides_seconds'] = round(time.perf_counter() - stage, 3)

stage = time.perf_counter()
rate = 24000
t = np.arange(rate * 8) / rate
notes = [261.63, 329.63, 392, 523.25, 392, 329.63, 293.66, 261.63]
frequency = np.repeat(notes, rate)
envelope = np.sin(np.pi * (t % 1))**2
audio = .22 * envelope * (np.sin(2*np.pi*frequency*t) + .25*np.sin(4*np.pi*frequency*t))
with wave.open(str(out / 'melody.wav'), 'wb') as sound:
    sound.setnchannels(1); sound.setsampwidth(2); sound.setframerate(rate)
    sound.writeframes((audio * 32767).astype(np.int16).tobytes())
subprocess.run(['espeak-ng', '-w', str(out / 'narration.wav'), 'This report was generated in an isolated offline Jarvis sandbox.'], check=True, timeout=15)
for name in ('melody.wav', 'narration.wav'):
    with wave.open(str(out / name)) as sound:
        assert sound.getnframes() > sound.getframerate()
subprocess.run(['ffmpeg', '-nostdin', '-n', '-v', 'error', '-loop', '1', '-i', str(out / 'chart.png'),
    '-i', str(out / 'melody.wav'), '-t', '6', '-r', '24', '-c:v', 'libx264', '-preset', 'fast', '-crf', '28',
    '-threads', '2', '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-shortest', str(out / 'video.mp4')], check=True, timeout=35)
probe = json.loads(subprocess.check_output(['ffprobe', '-v', 'error', '-show_streams', '-show_format', '-of', 'json', str(out / 'video.mp4')]))
assert float(probe['format']['duration']) >= 5.9
video = next(stream for stream in probe['streams'] if stream['codec_type'] == 'video')
assert (video['width'], video['height']) == (1280, 720)
steps['speech_melody_video_seconds'] = round(time.perf_counter() - stage, 3)

report = {'rows': n, 'sum_squares': actual, 'sql_verified': True, 'held_out_accuracy': accuracy,
    'artifacts_verified': ['groups.csv', 'chart.png', 'report.pdf', 'slides.pptx', 'melody.wav', 'narration.wav', 'video.mp4'],
    'phases': steps, 'seconds': round(time.perf_counter() - started, 3),
    'python_peak_rss_mib': round(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024, 1),
    'memory_scope': 'Linux ru_maxrss for fixture Python process; excludes total host/container/child peak.',
    'media_scope': 'Chart/document generation, procedural instrumental melody, eSpeak narration and FFmpeg still-image video; no diffusion image/video or commercial music API.'}
(out / 'validation.json').write_text(json.dumps(report, indent=2))
print(json.dumps(report))
